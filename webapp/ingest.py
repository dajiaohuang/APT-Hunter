"""
Converts an APT-Hunter *_Report.xlsx workbook into a local SQLite cache so the
web dashboard can filter/sort/page large sheets without holding everything in
memory. Ingestion is fully generic (driven by whatever headers a sheet has) so
it works for the core report as well as optional sheets (LLM Analysis, O365
hunts, sigma detections, ...).
"""
import json
import os
import re
import sqlite3
import time
from datetime import datetime, timezone

import openpyxl

RAW_LOG_COLUMN = "Original Event Log"
SUMMARY_SHEET_NAME = "Result Summary"
NUMERIC_RE = re.compile(r'^-?\d+(\.\d+)?$')
BATCH_SIZE = 2000
NUMERIC_SAMPLE_SIZE = 500

INCIDENT_TABLE = "_incident_timeline"
INCIDENT_TABLE_SQL = f'''
    CREATE TABLE IF NOT EXISTS {INCIDENT_TABLE} (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        sheet_name TEXT,
        source_row_id INTEGER,
        timestamp REAL,
        date_text TEXT,
        severity TEXT,
        description TEXT,
        computer TEXT,
        fields_json TEXT,
        note TEXT DEFAULT '',
        added_at REAL,
        parent_id INTEGER
    )
'''


TRIAGE_RUN_TABLE = "_triage_runs"
TRIAGE_RUN_TABLE_SQL = f'''
    CREATE TABLE IF NOT EXISTS {TRIAGE_RUN_TABLE} (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        started_at REAL,
        finished_at REAL,
        status TEXT,
        scope_json TEXT,
        stats_json TEXT,
        error TEXT
    )
'''

TRIAGE_FINDING_TABLE = "_triage_findings"
TRIAGE_FINDING_TABLE_SQL = f'''
    CREATE TABLE IF NOT EXISTS {TRIAGE_FINDING_TABLE} (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        run_id INTEGER,
        cluster_key TEXT,
        title TEXT,
        verdict TEXT,
        score INTEGER,
        severity TEXT,
        mitre TEXT,
        narrative TEXT,
        first_ts REAL,
        last_ts REAL,
        computer TEXT,
        event_count INTEGER,
        evidence_json TEXT,
        status TEXT DEFAULT 'pending',
        decided_at REAL,
        incident_ids_json TEXT,
        examined_by_agent INTEGER DEFAULT 0,
        tool_trace_json TEXT
    )
'''

# Tables that hold analyst/agent work product and must survive a cache rebuild.
PRESERVED_TABLES = [
    (INCIDENT_TABLE, INCIDENT_TABLE_SQL),
    (TRIAGE_RUN_TABLE, TRIAGE_RUN_TABLE_SQL),
    (TRIAGE_FINDING_TABLE, TRIAGE_FINDING_TABLE_SQL),
]


# Columns added after the tables first shipped. CREATE TABLE IF NOT EXISTS leaves an
# existing table alone, so a cache built by an older version needs them added here.
_ADDED_COLUMNS = {
    # NULL parent_id means a top-level row. Triage evidence points at the finding's
    # summary row so the timeline can group it instead of interleaving everything.
    INCIDENT_TABLE: [
        ("parent_id", "INTEGER"),
    ],
    TRIAGE_FINDING_TABLE: [
        ("examined_by_agent", "INTEGER DEFAULT 0"),
        ("tool_trace_json", "TEXT"),
    ],
}


def ensure_work_tables(conn):
    for name, sql in PRESERVED_TABLES:
        conn.execute(sql)
        wanted = _ADDED_COLUMNS.get(name)
        if not wanted:
            continue
        have = {r[1] for r in conn.execute(f"PRAGMA table_info({name})")}
        for column, decl in wanted:
            if column not in have:
                conn.execute(f'ALTER TABLE {name} ADD COLUMN "{column}" {decl}')


def _extract_preserved_rows(db_path):
    """Reads back incident-timeline and triage rows from a pre-existing cache
    DB so they survive a report re-ingest/rebuild (the cache DB is otherwise
    dropped and recreated from scratch). Returns {table: [row dicts]}."""
    preserved = {name: [] for name, _sql in PRESERVED_TABLES}
    if not os.path.exists(db_path):
        return preserved
    try:
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
    except sqlite3.Error:
        return preserved
    for name, _sql in PRESERVED_TABLES:
        try:
            preserved[name] = [dict(r) for r in conn.execute(f"SELECT * FROM {name}").fetchall()]
        except sqlite3.OperationalError:
            pass
    conn.close()
    return preserved


def _restore_preserved_rows(conn, preserved):
    ensure_work_tables(conn)
    for name, _sql in PRESERVED_TABLES:
        rows = preserved.get(name) or []
        if not rows:
            continue
        cols = list(rows[0].keys())
        col_list = ", ".join(f'"{c}"' for c in cols)
        placeholders = ", ".join("?" for _ in cols)
        conn.executemany(
            f"INSERT INTO {name} ({col_list}) VALUES ({placeholders})",
            [tuple(r.get(c) for c in cols) for r in rows],
        )


def cache_path_for(xlsx_path):
    base = os.path.splitext(xlsx_path)[0]
    return base + ".aptdash.db"


def needs_rebuild(xlsx_path, db_path):
    if not os.path.exists(db_path):
        return True
    return os.path.getmtime(xlsx_path) > os.path.getmtime(db_path)


def _safe_table_name(sheet_name):
    return "sheet_" + re.sub(r'[^0-9a-zA-Z_]', '_', sheet_name)


def _infer_columns(header_row):
    columns = []
    for i, h in enumerate(header_row):
        name = str(h).strip() if h is not None else ""
        if not name:
            name = f"column_{i + 1}"
        # Guard against duplicate headers colliding as SQL columns.
        base, n = name, 2
        while name in columns:
            name = f"{base}_{n}"
            n += 1
        columns.append(name)
    return columns


def _parse_result_summary(ws):
    """The Result Summary sheet is two stacked tables: Severity/Counts then,
    a few blank rows later, Detection/Counts. Returns (severity_counts dict,
    detections list of (name, count) sorted desc as they appear)."""
    severity_counts = {}
    detections = []
    mode = None
    for row in ws.iter_rows(values_only=True):
        if not row or all(c is None for c in row):
            continue
        a = row[0]
        b = row[1] if len(row) > 1 else None
        if a == "Severity" and b == "Counts":
            mode = "severity"
            continue
        if a == "Detection" and b == "Counts":
            mode = "detection"
            continue
        if a is None:
            continue
        try:
            count = int(b) if b is not None else 0
        except (TypeError, ValueError):
            count = 0
        if mode == "severity":
            severity_counts[str(a)] = count
        elif mode == "detection":
            detections.append([str(a), count])
    return severity_counts, detections


def _detect_numeric_columns(conn, table, columns):
    numeric = []
    for c in columns:
        cur = conn.execute(
            f'SELECT "{c}" FROM "{table}" WHERE "{c}" IS NOT NULL LIMIT ?',
            (NUMERIC_SAMPLE_SIZE,),
        )
        vals = [r[0] for r in cur.fetchall()]
        if vals and all(NUMERIC_RE.match(v) for v in vals):
            numeric.append(c)
    return numeric


def build_db(xlsx_path, db_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    preserved = _extract_preserved_rows(db_path)

    if os.path.exists(db_path):
        os.remove(db_path)
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=OFF")

    sheets_meta = []
    dashboard = {
        "severity_counts": {},
        "top_detections": [],
        "timeline": {},
        "top_computers": {},
        "top_event_ids": {},
    }

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        if sheet_name == SUMMARY_SHEET_NAME:
            severity_counts, detections = _parse_result_summary(ws)
            dashboard["severity_counts"] = severity_counts
            dashboard["top_detections"] = detections
            continue

        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            continue
        columns = _infer_columns(header)
        table = _safe_table_name(sheet_name)

        col_defs = ", ".join(f'"{c}" TEXT' for c in columns)
        conn.execute(f'CREATE TABLE "{table}" (_id INTEGER PRIMARY KEY AUTOINCREMENT, {col_defs})')
        col_list = ", ".join(f'"{c}"' for c in columns)
        placeholders = ", ".join("?" for _ in columns)
        insert_sql = f'INSERT INTO "{table}" ({col_list}) VALUES ({placeholders})'

        has_severity = "Severity" in columns
        has_timestamp = "timestamp" in columns
        has_computer = "Computer Name" in columns
        has_event_id = "Event ID" in columns
        sev_idx = columns.index("Severity") if has_severity else None
        ts_idx = columns.index("timestamp") if has_timestamp else None
        computer_idx = columns.index("Computer Name") if has_computer else None
        eid_idx = columns.index("Event ID") if has_event_id else None

        row_count = 0
        buffer = []
        for row in rows_iter:
            if row is None or all(c is None for c in row):
                continue
            vals = list(row[:len(columns)]) + [None] * (len(columns) - len(row))
            vals = [(str(v) if v is not None else None) for v in vals]
            buffer.append(vals)
            row_count += 1

            if has_computer and vals[computer_idx]:
                cn = vals[computer_idx]
                dashboard["top_computers"][cn] = dashboard["top_computers"].get(cn, 0) + 1
            if has_event_id and vals[eid_idx]:
                eid = vals[eid_idx]
                dashboard["top_event_ids"][eid] = dashboard["top_event_ids"].get(eid, 0) + 1
            if has_timestamp and vals[ts_idx]:
                try:
                    day = datetime.fromtimestamp(float(vals[ts_idx]), tz=timezone.utc).strftime("%Y-%m-%d")
                    sev = vals[sev_idx] if has_severity and vals[sev_idx] else "Unknown"
                    bucket = dashboard["timeline"].setdefault(day, {})
                    bucket[sev] = bucket.get(sev, 0) + 1
                except (ValueError, OSError, OverflowError):
                    pass

            if len(buffer) >= BATCH_SIZE:
                conn.executemany(insert_sql, buffer)
                buffer.clear()
        if buffer:
            conn.executemany(insert_sql, buffer)

        numeric_columns = _detect_numeric_columns(conn, table, columns)

        if has_severity:
            conn.execute(f'CREATE INDEX "idx_{table}_severity" ON "{table}"("Severity")')
        if has_timestamp:
            conn.execute(f'CREATE INDEX "idx_{table}_ts" ON "{table}"("timestamp")')
        if has_computer:
            conn.execute(f'CREATE INDEX "idx_{table}_computer" ON "{table}"("Computer Name")')

        kind = "event" if (has_severity and has_timestamp) else "summary"
        sheets_meta.append({
            "name": sheet_name,
            "table": table,
            "kind": kind,
            "columns": columns,
            "numeric_columns": numeric_columns,
            "row_count": row_count,
            "has_raw_log": RAW_LOG_COLUMN in columns,
        })

    conn.execute("CREATE TABLE _meta (key TEXT PRIMARY KEY, value TEXT)")
    conn.execute("INSERT INTO _meta VALUES (?, ?)", ("sheets", json.dumps(sheets_meta)))
    conn.execute("INSERT INTO _meta VALUES (?, ?)", ("dashboard", json.dumps(dashboard)))
    conn.execute("INSERT INTO _meta VALUES (?, ?)", (
        "source",
        json.dumps({"path": os.path.abspath(xlsx_path), "built_at": time.time()}),
    ))
    _restore_preserved_rows(conn, preserved)
    conn.commit()
    conn.close()
    wb.close()
    return db_path
